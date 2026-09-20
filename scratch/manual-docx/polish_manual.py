from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return False
    full = paragraph.text.replace(old, new)
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = full
    else:
        paragraph.add_run(full)
    return True


def style_document(input_path: Path, output_path: Path):
    doc = Document(str(input_path))
    props = doc.core_properties
    props.title = "Panduan Lengkap Penggunaan Sistem Keuangan Baru Pondok Pesantren Sukahideng"
    props.subject = "Manual operasional Sistem Keuangan Baru"
    props.author = "Pondok Pesantren Sukahideng"
    props.keywords = "keuangan pesantren, koperasi, pembayaran, uang jajan, rekonsiliasi"
    props.comments = "Versi 1.0"

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(27)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(12)

    if "Subtitle" in styles:
        subtitle = styles["Subtitle"]
        subtitle.font.name = "Aptos"
        subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        subtitle.font.size = Pt(13)
        subtitle.font.color.rgb = RGBColor(70, 70, 70)

    for name, size, before, after in (
        ("Heading 1", 18, 18, 8),
        ("Heading 2", 14, 14, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    # Repair cover and remove the old instruction that asked the reader to update fields.
    cover_map = {
        "PONDOK PESANTREN SUKAHIDENG": ("Pondok Pesantren Sukahideng", "Subtitle"),
        "PANDUAN LENGKAP PENGGUNAAN": ("Panduan Lengkap Penggunaan Sistem Keuangan Baru", "Title"),
        "SISTEM KEUANGAN BARU": ("Manual Operasional", "Subtitle"),
    }
    mapped = set()
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in cover_map and text not in mapped:
            new_text, style_name = cover_map[text]
            paragraph.text = new_text
            paragraph.style = styles[style_name]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            mapped.add(text)
        replace_text(
            paragraph,
            "Setelah dokumen dibuka, klik kanan pada daftar isi lalu pilih Update Field dan Update entire table agar nomor halaman menyesuaikan.",
            "Daftar isi dan nomor halaman telah diperbarui pada saat dokumen diterbitkan.",
        )
        replace_text(paragraph, "Klik kanan daftar isi ini, lalu pilih Update Field.", "Daftar isi")
        # Repair common UTF-8-to-Windows mojibake if the RTF importer encountered it.
        repairs = {
            "â–¡": "□",
            "â†’": "->",
            "â€“": "-",
            "â€”": "-",
            "Â": "",
        }
        for old, new in repairs.items():
            replace_text(paragraph, old, new)

    # Use real bullet styles for simple hyphen lists while preserving numbered procedures.
    for paragraph in doc.paragraphs:
        stripped = paragraph.text.strip()
        if stripped.startswith("- "):
            paragraph.text = stripped[2:]
            paragraph.style = styles["List Bullet"]
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True

    # Tables use a consistent dark-blue header, pale alternating rows, light borders, and breathing room.
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table)
        if table.rows:
            repeat_header(table.rows[0])
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                if r_idx == 0:
                    set_cell_shading(cell, "1F4E78")
                elif r_idx % 2 == 0:
                    set_cell_shading(cell, "F4F8FB")
                else:
                    set_cell_shading(cell, "FFFFFF")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_font(
                            run,
                            "Aptos",
                            size=8.5 if len(table.columns) >= 4 else 9.0,
                            bold=True if r_idx == 0 else None,
                            color=(255, 255, 255) if r_idx == 0 else (0, 0, 0),
                        )

    # Footer text remains subtle and readable.
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_font(run, "Aptos", size=8, color=(100, 100, 100))

    doc.save(str(output_path))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: polish_manual.py input.docx output.docx")
    style_document(Path(sys.argv[1]), Path(sys.argv[2]))
